# exportar_excel.py

import pandas as pd
from sqlalchemy import text
from database import SessionLocal
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def aplicar_formato_condicional(nombre_archivo: str):
    """
    Abre un archivo de Excel existente y aplica formato condicional a las celdas
    basado en el contenido de las columnas 'detalle' y 'saldo'.
    """
    try:
        # Carga el libro de trabajo y selecciona la hoja activa
        workbook = load_workbook(nombre_archivo)
        sheet = workbook.active

        # Define el estilo de relleno rojo
        relleno_rojo = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        # Busca los índices de las columnas 'saldo' y 'detalle'
        # Esto hace que el script no dependa del orden de las columnas
        encabezados = [cell.value for cell in sheet[1]]
        try:
            col_saldo_idx = encabezados.index('saldo') + 1
            col_detalle_idx = encabezados.index('detalle') + 1
        except ValueError as e:
            print(f"Error: No se encontró la columna requerida en el Excel: {e}")
            return

        # Itera sobre todas las filas (empezando desde la fila 2 para saltar el encabezado)
        for fila in sheet.iter_rows(min_row=2):
            celda_detalle = fila[col_detalle_idx - 1]
            celda_saldo = fila[col_saldo_idx - 1]

            # Condición 1: Si el detalle es "Saldo expirado"
            if celda_detalle.value == "Saldo expirado":
                # Pinta toda la fila de rojo
                for celda in fila:
                    celda.fill = relleno_rojo
                continue # Pasa a la siguiente fila para no aplicar la segunda regla

            # Condición 2: Si el saldo es menor a 50
            # Usamos 'try-except' por si la celda de saldo está vacía o no es un número
            try:
                if celda_saldo.value is not None and float(celda_saldo.value) < 50.0:
                    # Pinta toda la fila de rojo
                    for celda in fila:
                        celda.fill = relleno_rojo
            except (ValueError, TypeError):
                # Si la celda no contiene un número válido, simplemente la ignoramos
                continue

        # Guarda los cambios en el archivo
        workbook.save(nombre_archivo)
        print("🎨 Se ha aplicado el formato condicional con éxito.")

    except FileNotFoundError:
        print(f"Error: No se pudo encontrar el archivo {nombre_archivo} para aplicar el formato.")
    except Exception as e:
        print(f"❌ Ocurrió un error al aplicar el formato: {e}")


def exportar_consultas_a_excel(fecha_str: str):
    """
    Consulta los registros de la tabla 'unidad' para una fecha específica
    y los exporta a un archivo de Excel.
    """
    try:
        fecha_obj = datetime.strptime(fecha_str, '%Y-%m-%d').date()
    except ValueError:
        print("❌ Error: El formato de la fecha es incorrecto. Usa AAAA-MM-DD.")
        return

    db = SessionLocal()
    try:
        consulta_sql = text("""
            SELECT unidad, numero_celular, saldo, fecha_vencimiento, 
                   fecha_consulta, estado_consulta, detalle 
            FROM unidad 
            WHERE CAST(fecha_consulta AS DATE) = :fecha_a_buscar
        """)
        
        df = pd.read_sql(consulta_sql, db.bind, params={"fecha_a_buscar": fecha_obj})

        if df.empty:
            print(f"⚠️ No se encontraron registros para la fecha {fecha_str}.")
            return

        print(f"✅ Se encontraron {len(df)} registros. Exportando a Excel...")
        
        nombre_archivo = f"Reporte_Saldos_{fecha_str}.xlsx"
        df.to_excel(nombre_archivo, index=False, sheet_name='Reporte')
        
        print(f"¡Éxito! Los datos han sido guardados en: {nombre_archivo}")

        # --- ¡LLAMADA A LA NUEVA FUNCIÓN! ---
        # Después de crear el archivo, le aplicamos el formato.
        aplicar_formato_condicional(nombre_archivo)

    except Exception as e:
        print(f"❌ Ocurrió un error durante la exportación: {e}")
    finally:
        db.close()

# --- EJECUCIÓN DEL SCRIPT ---
if __name__ == "__main__":
    fecha_input = input("Ingresa la fecha para el reporte (formato AAAA-MM-DD): ")
    exportar_consultas_a_excel(fecha_input)